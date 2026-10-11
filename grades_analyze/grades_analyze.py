import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from docx import Document
import os



try:
    # 读取grades和syllabus的excel
    df_grades = pd.read_excel('grades.xlsx')
    df_syllabus = pd.read_excel('syllabus.xlsx')

    # 横变纵
    df_melt_grades = df_grades.melt(id_vars='Assignment', var_name='Name', value_name='Score')


    # 合并两个表
    combined = pd.merge(df_melt_grades, df_syllabus, on='Assignment')

    total_final = 0
    count = 0

    # 生成word
    doc = Document()
    doc.add_heading('Student Performance Analysis Report', 0)
    doc.add_paragraph('This report provides an automated analysis of student academic behavior and final weighted scores.')

    # 开始数据处理
    for name, group in combined.groupby('Name'):
        # 计算总分
        final_score = (group['Score'] * group['Weight']).sum()
        total_final = final_score + total_final
        count = count + 1

        # 计算k（x: 作业，y: 成绩）和方差
        x = np.arange(len(group))
        y = group['Score'].values
        k, b = np.polyfit(x, y, 1)
        v = np.var(y)

        # 判定等地
        if final_score >= 93:
            Grade = 'A'
            GPA = '4.0'
        elif 88 <= final_score < 93:
            Grade = 'AB'
            GPA = '3.5'
        elif 83 <= final_score < 88:
            Grade = 'B'
            GPA = '3.0'
        elif 78 <= final_score < 83:
            Grade = 'BC'
            GPA = '2.5'
        elif 70 <= final_score < 78:
            Grade = 'C'
            GPA = '2.0'
        else:
            Grade = 'N/A'
            GPA = 'N/A'

        # 判定学习状态
        if k > 2:
            behavior = "Significant Improvement"
        elif k < -2:
            behavior = "Regression Observed"
        else:
            behavior = "Stable Tendency"

        if v > 20:
            performance = "Unstable Performance"
        else:
            performance = "Stable Performance"

        # 生成折线图
        plt.figure(figsize=(6, 4))
        plt.plot(group['Assignment'], y, marker='o', color='blue', label='Actual Score')
        plt.plot(x, k * x + b, color='red', linestyle='--', label='Trend Line')
        plt.title(f'Academic Performance: {name}')
        plt.xlabel('Assignments')
        plt.ylabel('Score')
        plt.legend()

        png_name = f"{name}_trend.png"
        plt.savefig(png_name)
        plt.close()

        # 生成个人报告
        doc.add_heading(f'Student: {name}', level=1)
        doc.add_paragraph(f'Final Weighted Score: {final_score:.2f}')
        doc.add_paragraph(f'Grade: {Grade}')
        doc.add_paragraph(f'GPA: {GPA}')
        doc.add_paragraph(f'Trend Slope (k): {k:.2f}')
        doc.add_paragraph(f'Academic Behavior: {behavior}')
        doc.add_paragraph(f'Student Performance: {performance}')
        doc.add_picture(png_name)
        doc.add_page_break()

        # 删除ssd中的图片
        os.remove(png_name)

    # 保存文档
    doc_name = 'Semester_Report.docx'
    doc.save(doc_name)

    # 计算平均分
    avg_score = total_final / count

    # 打印结果
    print("\n" + "=" * 40)
    print("FINAL RESULT")
    print(f'Average Score: {avg_score:.2f}')
    print("=" * 40)
    print("\nSUCCESS: 'Semester_Report.docx' has been generated!")

    # 打印失败结果
except FileNotFoundError:
    print("\n" + "=" * 40)
    print("FINAL RESULT")
    print("=" * 40)
    print("ERROR: Please ensure 'grades.xlsx' and 'syllabus.xlsx' exist in the folder.")
